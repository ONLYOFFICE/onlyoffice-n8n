"""Regenerate the office files used by the automation tests.

Requires: pip install python-docx openpyxl python-pptx msoffcrypto-tool
Run from this folder: python generate.py
"""

import io

import docx
import msoffcrypto.format.ooxml
import openpyxl
import pptx

PASSWORD = "test-password"


def sample_docx():
    document = docx.Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("The quick brown fox jumps over the lazy dog.")
    document.add_heading("Details", level=2)
    document.add_paragraph("Second section text.")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Name", "Score"), ("Alice", "95"), ("Bob", "87")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    return document


def template_docx():
    document = docx.Document()
    document.add_paragraph("Hello {{name}}, today is {{date}}.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Role"
    table.rows[0].cells[1].text = "{{role}}"
    return document


def sample_xlsx():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in [("Name", "Score", "Status"), ("Alice", 95, "Open"), ("Bob", 87, "Open")]:
        sheet.append(row)
    return workbook


def sample_pptx():
    presentation = pptx.Presentation()
    for title in ("First slide", "Second slide"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"{title} body"
    return presentation


def main():
    sample_docx().save("sample.docx")
    template_docx().save("template.docx")
    sample_xlsx().save("sample.xlsx")
    sample_pptx().save("sample.pptx")

    plain = io.BytesIO()
    sample_docx().save(plain)
    plain.seek(0)
    with open("protected.docx", "wb") as target:
        msoffcrypto.format.ooxml.OOXMLFile(plain).encrypt(PASSWORD, target)


if __name__ == "__main__":
    main()
